"""Build the Blog schedules base sheet from an already generated workbook."""
from __future__ import annotations

from datetime import datetime, time, timedelta
import re
from typing import Any

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

BLOG_SHEET_NAME = "Base Horarios para el Blog"
ACTIVITIES_SHEET_NAME = "Actividades"
BLOG_COLUMNS = [
    "ficha",
    "programa",
    "jornada",
    "dia",
    "hora_inicio",
    "hora_fin",
    "competencia",
    "instructor",
    "aula",
    "fecha_inicio",
    "fecha_fin",
]
VALID_JOURNEYS = {"DIURNA", "NOCTURNA", "MIXTA", "MADRUGADA", "FIN DE SEMANA"}


def generar_base_horarios_blog(workbook: Workbook) -> Worksheet:
    """Add the Blog schedules base sheet using the existing Actividades sheet.

    This function is intentionally additive: it only deletes/recreates
    ``Base Horarios para el Blog`` and leaves every other sheet untouched.
    """
    if ACTIVITIES_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f'El workbook no contiene la hoja "{ACTIVITIES_SHEET_NAME}"')
    if BLOG_SHEET_NAME in workbook.sheetnames:
        del workbook[BLOG_SHEET_NAME]
    worksheet = workbook.create_sheet(BLOG_SHEET_NAME)
    worksheet.append(BLOG_COLUMNS)

    activities_sheet = workbook[ACTIVITIES_SHEET_NAME]
    headers = _header_map(activities_sheet)
    for source_row in activities_sheet.iter_rows(min_row=2, values_only=True):
        day = _cell(source_row, headers, "Día")
        hour_value = _cell(source_row, headers, "Hora")
        if not day or not hour_value:
            continue
        start_time = _parse_time(hour_value)
        if start_time is None:
            continue
        duration = _parse_duration(_cell(source_row, headers, "Duración"))
        end_time = _add_hours(start_time, duration)
        ficha, programa, jornada = _split_group(str(_cell(source_row, headers, "Grupo") or ""))
        instructor = str(_cell(source_row, headers, "Docente") or "").strip()
        if instructor == "SIN DOCENTE":
            instructor = ""
        worksheet.append([
            ficha,
            programa,
            jornada,
            day,
            start_time,
            end_time,
            _cell(source_row, headers, "Materia") or "",
            instructor,
            _cell(source_row, headers, "Salón") or "",
            "",
            "",
        ])

    for row in worksheet.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            if cell.value:
                cell.number_format = "h:mm"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    _auto_width(worksheet)
    return worksheet


def _header_map(worksheet: Worksheet) -> dict[str, int]:
    return {str(cell.value).strip(): index for index, cell in enumerate(worksheet[1]) if cell.value is not None}


def _cell(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    index = headers.get(name)
    return row[index] if index is not None and index < len(row) else None


def _split_group(value: str) -> tuple[str, str, str]:
    parts = [part.strip() for part in value.split(" - ") if part and part.strip()]
    if not parts:
        return "", "", ""
    raw_ficha = parts[0]
    digit_match = re.match(r"^\s*(\d+)", raw_ficha)
    ficha = digit_match.group(1) if digit_match else raw_ficha
    raw_journey = parts[-1].upper()
    jornada = raw_journey if raw_journey in VALID_JOURNEYS else ""
    program_parts = parts[1:-1] if jornada else parts[1:]
    programa = " - ".join(program_parts).strip().rstrip(" .")
    return ficha, programa, jornada


def _parse_time(value: Any) -> time | None:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    raw_value = str(value or "").strip()
    match = re.search(r"\d{1,2}:\d{2}", raw_value)
    if match:
        raw_value = match.group(0)
    for fmt in ("%H:%M", "%I:%M %p"):
        try:
            return datetime.strptime(raw_value, fmt).time()
        except ValueError:
            continue
    return None


def _parse_duration(value: Any) -> int:
    try:
        return max(int(value), 1)
    except (TypeError, ValueError):
        return 1


def _add_hours(value: time, hours: int) -> time:
    base = datetime.combine(datetime.today(), value)
    return (base + timedelta(hours=hours)).time()


def _auto_width(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        max_len = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        worksheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 50)
