"""Reusable Excel formatting helpers."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="EAF2F8")
BORDER = Border(*(Side(style="thin", color="B7B7B7"),) * 4)


def style_table(ws: Worksheet, header_row: int = 1) -> None:
    """Apply a professional table style to a worksheet."""
    max_row = ws.max_row
    max_col = ws.max_column
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for row in ws.iter_rows(min_row=header_row + 1, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = ALT_FILL
    if max_row >= header_row and max_col:
        ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    auto_width(ws)


def auto_width(ws: Worksheet) -> None:
    """Resize columns using visible content length."""
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        max_len = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 60)
