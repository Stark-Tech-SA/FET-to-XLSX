"""Excel exporter for parsed FET data."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import re
from typing import Iterable, NamedTuple

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from fet_to_xlsx.exporter.formatting import style_table
from fet_to_xlsx.parser.models import Activity, FetData, ScheduledActivity


class GridResult(NamedTuple):
    """Built timetable grid for a single teacher/group/room resource."""

    grid: dict[tuple[str, int], Activity]
    spans: dict[tuple[str, int], int]
    conflicts: list[tuple[str, Activity, Activity]]


class ExcelExporter:
    """Create a formatted .xlsx workbook from typed FET data."""

    WEEKDAYS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]

    def export(self, data: FetData, output_path: str | Path) -> Path:
        path = Path(output_path)
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        wb = Workbook()
        conflicts: list[tuple[str, str, Activity, Activity]] = []
        self._general(wb.active, data)
        self._teachers(wb.create_sheet("Docentes"), data)
        self._subjects(wb.create_sheet("Materias"), data)
        self._groups(wb.create_sheet("Grupos"), data)
        self._rooms(wb.create_sheet("Salones"), data)
        self._activities(wb.create_sheet("Actividades"), data)
        self._constraints(wb.create_sheet("Restricciones"), data)
        self._unscheduled(wb.create_sheet("Sin horario"), data)
        conflicts.extend(self._resource_timetables(wb, data, "group"))
        conflicts.extend(self._resource_timetables(wb, data, "teacher"))
        conflicts.extend(self._resource_timetables(wb, data, "room"))
        self._conflicts(wb.create_sheet("Conflictos"), conflicts)
        wb.save(path)
        return path

    def _general(self, ws: Worksheet, data: FetData) -> None:
        ws.title = "Información General"
        rows = [
            ("Campo", "Valor"),
            ("Institución", data.institution.name),
            ("Versión", data.institution.version),
            ("Fecha de exportación", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Número de docentes", len(data.teachers)),
            ("Número de estudiantes", sum(g.number_of_students or 0 for g in data.groups)),
            ("Número de grupos", len(data.groups)),
            ("Número de actividades", len(data.activities)),
            ("Número de materias", len(data.subjects)),
            ("Número de salones", len(data.rooms)),
            ("Número de restricciones", len(data.constraints)),
            ("Actividades sin horario", len(data.unscheduled_activity_ids)),
        ]
        for row in rows:
            ws.append(row)
        style_table(ws)

    def _teachers(self, ws: Worksheet, data: FetData) -> None:
        ws.append(["Nombre", "Comentarios", "Disponibilidad"])
        for teacher in data.teachers:
            ws.append([teacher.name, teacher.comments, self._stringify(teacher.availability)])
        style_table(ws)

    def _subjects(self, ws: Worksheet, data: FetData) -> None:
        ws.append(["Nombre", "Comentarios"])
        for subject in data.subjects:
            ws.append([subject.name, subject.comments])
        style_table(ws)

    def _groups(self, ws: Worksheet, data: FetData) -> None:
        ws.append(["Año", "Grupo", "Subgrupo", "Número de estudiantes"])
        for group in data.groups:
            ws.append([group.year, group.group, group.subgroup, group.number_of_students])
        style_table(ws)

    def _rooms(self, ws: Worksheet, data: FetData) -> None:
        ws.append(["Nombre", "Capacidad", "Edificio", "Comentarios"])
        for room in data.rooms:
            ws.append([room.name, room.capacity, room.building, room.comments])
        style_table(ws)

    def _activities(self, ws: Worksheet, data: FetData) -> None:
        placements = {placement.activity_id: placement for placement in data.scheduled_activities}
        ws.append([
            "ID", "Docente", "Materia", "Grupo", "Duración", "Cantidad de estudiantes",
            "Actividad activa", "Día", "Hora", "Salón", "Comentarios",
        ])
        for activity in data.activities:
            placement = placements.get(activity.id)
            ws.append([
                activity.id,
                ", ".join(activity.teachers),
                activity.subject,
                ", ".join(activity.students),
                activity.duration,
                activity.student_count,
                "Sí" if activity.active else "No",
                placement.day if placement else "",
                placement.hour if placement else "",
                placement.room if placement else "",
                activity.comments,
            ])
        style_table(ws)

    def _constraints(self, ws: Worksheet, data: FetData) -> None:
        ws.append(["Categoría", "Tipo", "Peso", "Campos"])
        for constraint in data.constraints:
            ws.append([constraint.category, constraint.type, constraint.weight, self._stringify(constraint.fields)])
        style_table(ws)

    def _unscheduled(self, ws: Worksheet, data: FetData) -> None:
        activities = {activity.id: activity for activity in data.activities}
        ws.append(["ID", "Docente", "Materia", "Grupo", "Duración", "Motivo"])
        for activity_id in data.unscheduled_activity_ids:
            activity = activities.get(activity_id)
            if activity is None:
                continue
            ws.append([
                activity.id,
                ", ".join(activity.teachers),
                activity.subject,
                ", ".join(activity.students),
                activity.duration,
                "No existe ConstraintActivityPreferredStartingTime para esta actividad",
            ])
        style_table(ws)

    def _conflicts(self, ws: Worksheet, conflicts: list[tuple[str, str, Activity, Activity]]) -> None:
        ws.append(["Vista", "Recurso", "Actividad en conflicto", "Materia", "Actividad existente", "Materia existente"])
        for mode, resource, activity, existing in conflicts:
            ws.append([self._base_title(mode), resource, activity.id, activity.subject, existing.id, existing.subject])
        style_table(ws)

    def _resource_timetables(self, wb: Workbook, data: FetData, mode: str) -> list[tuple[str, str, Activity, Activity]]:
        resources = self._resources(data, mode)
        conflicts: list[tuple[str, str, Activity, Activity]] = []
        if not resources:
            self._empty_timetable(wb.create_sheet(self._base_title(mode)), data)
            return conflicts
        for resource in resources:
            ws = wb.create_sheet(self._unique_title(wb, f"{self._prefix(mode)} {resource}"))
            result = self._timetable(ws, data, mode, resource)
            conflicts.extend((mode, resource, activity, existing) for _, activity, existing in result.conflicts)
        return conflicts

    def _empty_timetable(self, ws: Worksheet, data: FetData) -> None:
        ws.append(["Hora", *self._days(data)])
        ws.append(["Sin recursos definidos", *[""] * len(self._days(data))])
        style_table(ws)

    def _timetable(self, ws: Worksheet, data: FetData, mode: str, resource: str) -> GridResult:
        ws.append(["Hora", *self._days(data)])
        activities = {activity.id: activity for activity in data.activities}
        resource_activities = self._activities_for_resource(data, activities, mode, resource)
        result = self._build_grid(data, resource_activities)
        if not data.scheduled_activities:
            ws.append(["Sin solución de horario incluida", *[""] * len(self._days(data))])
            style_table(ws)
            return result
        for index, hour in enumerate(data.hours or ["Hora"]):
            row = [hour]
            for day in self._days(data):
                activity = result.grid.get((day, index))
                row.append(self._format_activity(activity, self._placement_for(data, activity), mode) if activity else "")
            ws.append(row)
        style_table(ws)
        self._merge_duration_cells(ws, result)
        return result

    def _build_grid(self, data: FetData, activities: list[Activity]) -> GridResult:
        hour_index = {hour: index for index, hour in enumerate(data.hours)}
        placements = {placement.activity_id: placement for placement in data.scheduled_activities}
        occupied: dict[str, dict[int, Activity]] = defaultdict(dict)
        grid: dict[tuple[str, int], Activity] = {}
        spans: dict[tuple[str, int], int] = {}
        conflicts: list[tuple[str, Activity, Activity]] = []
        for activity in activities:
            placement = placements.get(activity.id)
            if placement is None or placement.day not in self._days(data) or placement.hour not in hour_index:
                continue
            start = hour_index[placement.hour]
            duration = max(activity.duration or 1, 1)
            end = min(start + duration - 1, len(data.hours) - 1)
            span_indexes = range(start, end + 1)
            existing = next((occupied[placement.day][idx] for idx in span_indexes if idx in occupied[placement.day]), None)
            if existing is not None:
                conflicts.append((placement.day, activity, existing))
                continue
            for idx in span_indexes:
                occupied[placement.day][idx] = activity
            grid[(placement.day, start)] = activity
            spans[(placement.day, start)] = end - start + 1
        return GridResult(grid=grid, spans=spans, conflicts=conflicts)

    def _merge_duration_cells(self, ws: Worksheet, result: GridResult) -> None:
        for (day, start_index), span in result.spans.items():
            if span <= 1:
                continue
            day_column = self._days_from_header(ws).index(day) + 2
            start_row = start_index + 2
            ws.merge_cells(start_row=start_row, start_column=day_column, end_row=start_row + span - 1, end_column=day_column)

    def _activities_for_resource(
        self,
        data: FetData,
        activities: dict[str, Activity],
        mode: str,
        resource: str,
    ) -> list[Activity]:
        selected: list[Activity] = []
        for placement in data.scheduled_activities:
            activity = activities.get(placement.activity_id)
            if activity and self._matches(resource, placement, activity, mode):
                selected.append(activity)
        return selected

    def _placement_for(self, data: FetData, activity: Activity | None) -> ScheduledActivity | None:
        if activity is None:
            return None
        for placement in data.scheduled_activities:
            if placement.activity_id == activity.id:
                return placement
        return None

    def _format_activity(self, activity: Activity | None, placement: ScheduledActivity | None, mode: str) -> str:
        if activity is None:
            return ""
        lines = [activity.subject]
        if mode != "teacher" and activity.teachers:
            lines.append(", ".join(activity.teachers))
        if mode != "group" and activity.students:
            lines.append(", ".join(activity.students))
        if mode != "room" and placement and placement.room:
            lines.append(placement.room)
        return "\n".join(line for line in lines if line)

    def _matches(self, resource: str, placement: ScheduledActivity, activity: Activity, mode: str) -> bool:
        return (
            (mode == "teacher" and resource in activity.teachers)
            or (mode == "group" and resource in activity.students)
            or (mode == "room" and resource == placement.room)
        )

    def _resources(self, data: FetData, mode: str) -> list[str]:
        if mode == "teacher":
            return sorted({teacher for activity in data.activities for teacher in activity.teachers})
        if mode == "group":
            return sorted({student for activity in data.activities for student in activity.students})
        return [room.name for room in data.rooms]

    def _placements(self, scheduled: Iterable[ScheduledActivity]) -> dict[tuple[str, str], list[ScheduledActivity]]:
        placements: dict[tuple[str, str], list[ScheduledActivity]] = defaultdict(list)
        for placement in scheduled:
            placements[(placement.day, placement.hour)].append(placement)
        return placements

    def _days(self, data: FetData) -> list[str]:
        return data.days or self.WEEKDAYS

    def _days_from_header(self, ws: Worksheet) -> list[str]:
        return [str(cell.value) for cell in ws[1][1:]]

    def _base_title(self, mode: str) -> str:
        return {"teacher": "Horario Docente", "group": "Horario Grupo", "room": "Horario Salón"}[mode]

    def _prefix(self, mode: str) -> str:
        return {"teacher": "Docente", "group": "Grupo", "room": "Salón"}[mode]

    def _unique_title(self, wb: Workbook, title: str) -> str:
        safe_title = re.sub(r"[\\/*?:\[\]]", "-", title).strip() or "Horario"
        safe_title = safe_title[:31]
        if safe_title not in wb.sheetnames:
            return safe_title
        for index in range(1, 1000):
            suffix = f" {index}"
            candidate = f"{safe_title[:31 - len(suffix)]}{suffix}"
            if candidate not in wb.sheetnames:
                return candidate
        return safe_title[:28] + " 999"

    def _stringify(self, value: object) -> str:
        return str(value).replace("{", "").replace("}", "") if value else ""
