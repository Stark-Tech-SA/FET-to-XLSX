"""Exporter for the existing CMM database workbook format."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, time, timedelta
import logging
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from fet_to_xlsx.parser.fet_reader import FetReader
from fet_to_xlsx.parser.models import Activity, FetData, ScheduledActivity
from fet_to_xlsx.parser.parser import FetParser
from fet_to_xlsx.parser.solution_finder import SolutionFinder

LOGGER = logging.getLogger(__name__)

CMM_COLUMNS = [
    "ficha",
    "programa",
    "jornada",
    "dia",
    "hora_inicio",
    "hora_fin",
    "competencia",
    "instructor",
    "aula",
    "fecha_inicio",
    "fecha_fin",
]

VALID_JOURNEYS = {"DIURNA", "NOCTURNA", "MIXTA", "MADRUGADA", "FIN DE SEMANA"}
DEFAULT_TEACHER = "SIN DOCENTE"
DEFAULT_GROUP = "SIN GRUPO"


@dataclass(frozen=True, slots=True)
class StudentInfo:
    """Normalized fields extracted from a FET Students string."""

    ficha: str = ""
    programa: str = ""
    jornada: str = ""


def export_horarios_fet_sheet_from_fet(
    fet_path: str | Path,
    workbook_path: str | Path = "cmm-horarios-base-datos.xlsx",
    sheet_name: str = "Horarios_FET",
) -> Path:
    """Parse a .fet file and update the CMM workbook with the Horarios_FET sheet."""
    root = FetReader().read(fet_path)
    data = FetParser().parse(root)
    if not data.scheduled_activities:
        data.scheduled_activities = SolutionFinder().find_for(fet_path, data)
        scheduled_ids = {placement.activity_id for placement in data.scheduled_activities}
        data.unscheduled_activity_ids = [activity.id for activity in data.activities if activity.id not in scheduled_ids]
    return export_horarios_fet_sheet(data, workbook_path, sheet_name)


def export_horarios_fet_sheet(
    data: FetData,
    workbook_path: str | Path = "cmm-horarios-base-datos.xlsx",
    sheet_name: str = "Horarios_FET",
) -> Path:
    """Create/replace a Horarios_FET sheet inside an existing CMM workbook.

    The workbook is opened with ``load_workbook`` so existing sheets such as
    ``Horarios`` and ``Instrucciones`` keep their content and formatting.
    """
    path = Path(workbook_path)
    workbook = load_workbook(path)
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    _write_header(worksheet)
    rows = build_horarios_fet_rows(data)
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            if column_index in {5, 6} and value:
                cell.number_format = "h:mm"
    _auto_width(worksheet)
    workbook.save(path)
    return path


def build_horarios_fet_rows(data: FetData) -> list[list[object]]:
    """Build rows for the CMM Horarios_FET sheet from parsed FET data."""
    placements = {placement.activity_id: placement for placement in data.scheduled_activities}
    rows: list[list[object]] = []
    for activity in data.activities:
        placement = placements.get(activity.id)
        if placement is None or not placement.day or not placement.hour:
            LOGGER.warning("Actividad sin horario asignado omitida de Horarios_FET: %s", activity.id)
            continue
        hour_start = _parse_time(placement.hour)
        if hour_start is None:
            LOGGER.warning("Actividad con hora inválida omitida de Horarios_FET: %s (%s)", activity.id, placement.hour)
            continue
        hour_end = _calculate_end_time(data.hours, placement.hour, activity.duration or 1)
        instructor = _external_instructor(activity)
        for students in _external_students(activity):
            student_info = parse_students_field(students)
            rows.append([
                student_info.ficha,
                student_info.programa,
                student_info.jornada,
                placement.day,
                hour_start,
                hour_end,
                activity.subject,
                instructor,
                placement.room,
                "",
                "",
            ])
    return rows


def parse_students_field(value: str) -> StudentInfo:
    """Split a FET Students value into ficha, programa and jornada."""
    parts = [part.strip() for part in value.split(" - ") if part.strip()]
    if not parts:
        return StudentInfo()
    raw_journey = parts[-1].upper()
    jornada = raw_journey if raw_journey in VALID_JOURNEYS else ""
    middle_end = -1 if jornada else len(parts)
    raw_ficha = parts[0]
    digit_match = re.match(r"^\s*(\d+)", raw_ficha)
    ficha = digit_match.group(1) if digit_match else raw_ficha
    program_parts = parts[1:middle_end]
    programa = " - ".join(program_parts).strip().rstrip(" .")
    return StudentInfo(ficha=ficha, programa=programa, jornada=jornada)


def _external_students(activity: Activity) -> list[str]:
    return [student for student in activity.students if student and student != DEFAULT_GROUP] or [""]


def _external_instructor(activity: Activity) -> str:
    teachers = [teacher for teacher in activity.teachers if teacher and teacher != DEFAULT_TEACHER]
    return ", ".join(teachers)


def _calculate_end_time(hours: list[str], start_hour: str, duration: int) -> time:
    parsed_hours = [_parse_time(hour) for hour in hours]
    hour_index = {hour: index for index, hour in enumerate(hours)}
    start_index = hour_index.get(start_hour)
    if start_index is not None and start_index + duration < len(parsed_hours):
        candidate = parsed_hours[start_index + duration]
        if candidate is not None:
            return candidate
    fallback_base = next((hour for hour in reversed(parsed_hours) if hour is not None), _parse_time(start_hour))
    if fallback_base is None:
        return time(0, 0)
    step = _infer_hour_step(parsed_hours)
    return _add_timedelta(fallback_base, step)


def _infer_hour_step(hours: list[time | None]) -> timedelta:
    valid_hours = [hour for hour in hours if hour is not None]
    if len(valid_hours) >= 2:
        first = _as_datetime(valid_hours[0])
        second = _as_datetime(valid_hours[1])
        delta = second - first
        if delta.total_seconds() > 0:
            return delta
    return timedelta(hours=1)


def _parse_time(value: str) -> time | None:
    raw_value = (value or "").strip()
    if not raw_value:
        return None
    match = re.search(r"\d{1,2}:\d{2}", raw_value)
    if match:
        raw_value = match.group(0)
    for fmt in ("%H:%M", "%I:%M %p"):
        try:
            return datetime.strptime(raw_value, fmt).time()
        except ValueError:
            continue
    return None


def _add_timedelta(value: time, delta: timedelta) -> time:
    return (_as_datetime(value) + delta).time()


def _as_datetime(value: time) -> datetime:
    return datetime.combine(datetime.today(), value)


def _write_header(worksheet: Worksheet) -> None:
    worksheet.append(CMM_COLUMNS)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:K1"


def _auto_width(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        max_len = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        worksheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 50)
