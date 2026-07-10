import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import load_workbook

from fet_to_xlsx.exporter.excel_exporter import ExcelExporter
from fet_to_xlsx.parser.models import (
    Activity, FetData, Institution, Room, ScheduledActivity, StudentsGroup,
    Subject, Teacher,
)


def test_exporter_creates_expected_sheets(tmp_path):
    data = FetData(
        institution=Institution(name="Colegio Demo", version="6.0"),
        teachers=[Teacher("Ana")],
        subjects=[Subject("Matemáticas")],
        groups=[StudentsGroup(year="1", group="1A")],
        rooms=[Room("101")],
        activities=[Activity(id="1", teachers=["Ana"], subject="Matemáticas", students=["1A"], duration=1)],
        days=["Lunes"],
        hours=["08:00"],
        scheduled_activities=[ScheduledActivity(activity_id="1", day="Lunes", hour="08:00", room="101")],
    )
    output = ExcelExporter().export(data, tmp_path / "demo.xlsx")
    workbook = load_workbook(output)
    assert "Información General" in workbook.sheetnames
    assert "Docentes" in workbook.sheetnames
    assert "Restricciones" in workbook.sheetnames
    assert "Grupo 1A" in workbook.sheetnames
    assert workbook["Grupo 1A"]["B2"].value == "Matemáticas\nAna\n101"


def test_exporter_merges_duration_and_reports_conflicts(tmp_path):
    data = FetData(
        activities=[
            Activity(id="1", teachers=["Ana"], subject="Bloque 1", students=["1A"], duration=2),
            Activity(id="2", teachers=["Ana"], subject="Bloque 2", students=["1A"], duration=1),
        ],
        days=["Lunes"],
        hours=["6:00", "7:00", "8:00"],
        scheduled_activities=[
            ScheduledActivity(activity_id="1", day="Lunes", hour="6:00"),
            ScheduledActivity(activity_id="2", day="Lunes", hour="7:00"),
        ],
    )
    output = ExcelExporter().export(data, tmp_path / "conflicts.xlsx")
    workbook = load_workbook(output)
    assert "B2:B3" in [str(range_) for range_ in workbook["Grupo 1A"].merged_cells.ranges]
    assert workbook["Grupo 1A"]["B2"].value == "Bloque 1\nAna"
    assert workbook["Conflictos"].max_row == 3  # conflicto en vista grupo y docente
