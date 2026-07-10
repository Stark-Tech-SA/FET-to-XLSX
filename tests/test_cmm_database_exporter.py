import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook, load_workbook

from fet_to_xlsx.exporter.cmm_database_exporter import (
    CMM_COLUMNS,
    build_horarios_fet_rows,
    export_horarios_fet_sheet,
    parse_students_field,
)
from fet_to_xlsx.parser.models import Activity, FetData, ScheduledActivity


def test_parse_students_field_extracts_ficha_programa_jornada():
    info = parse_students_field("2925097 - MANTENIMIENTO ELECTROMECANICO INDUSTRIAL . - MIXTA")
    assert info.ficha == "2925097"
    assert info.programa == "MANTENIMIENTO ELECTROMECANICO INDUSTRIAL"
    assert info.jornada == "MIXTA"


def test_parse_students_field_keeps_non_numeric_ficha_and_unknown_jornada():
    info = parse_students_field("Tco Mecánico 1 - Programa Especial - TARDE")
    assert info.ficha == "Tco Mecánico 1"
    assert info.programa == "Programa Especial - TARDE"
    assert info.jornada == ""


def test_build_rows_skips_unscheduled_and_uses_time_values():
    data = FetData(
        activities=[
            Activity(
                id="1",
                teachers=["Ana"],
                subject="Competencia completa",
                students=["2925097 - MANTENIMIENTO ELECTROMECANICO INDUSTRIAL . - MIXTA"],
                duration=4,
            ),
            Activity(id="2", teachers=["SIN DOCENTE"], subject="Sin ubicar", students=["SIN GRUPO"], duration=1),
        ],
        hours=["18:00", "19:00", "20:00", "21:00"],
        scheduled_activities=[ScheduledActivity(activity_id="1", day="Miércoles", hour="18:00", room="Aula 1")],
    )
    rows = build_horarios_fet_rows(data)
    assert len(rows) == 1
    assert rows[0][0:4] == ["2925097", "MANTENIMIENTO ELECTROMECANICO INDUSTRIAL", "MIXTA", "Miércoles"]
    assert rows[0][4].hour == 18
    assert rows[0][5].hour == 22
    assert rows[0][6:9] == ["Competencia completa", "Ana", "Aula 1"]
    assert rows[0][9:11] == ["", ""]


def test_export_horarios_fet_sheet_preserves_existing_sheets_and_replaces_target(tmp_path):
    path = tmp_path / "cmm-horarios-base-datos.xlsx"
    workbook = Workbook()
    workbook.active.title = "Horarios"
    workbook.active["A1"] = "mantener"
    workbook.create_sheet("Instrucciones")["A1"] = "también mantener"
    workbook.create_sheet("Horarios_FET")["A1"] = "anterior"
    workbook.save(path)
    data = FetData(
        activities=[Activity(id="1", teachers=["Ana"], subject="Competencia", students=["123 - Prog - DIURNA"], duration=1)],
        hours=["6:00", "7:00"],
        scheduled_activities=[ScheduledActivity(activity_id="1", day="Lunes", hour="6:00", room="101")],
    )

    export_horarios_fet_sheet(data, path)

    updated = load_workbook(path)
    assert updated["Horarios"]["A1"].value == "mantener"
    assert updated["Instrucciones"]["A1"].value == "también mantener"
    assert [cell.value for cell in updated["Horarios_FET"][1]] == CMM_COLUMNS
    assert updated["Horarios_FET"]["A2"].value == "123"
    assert updated["Horarios_FET"]["E2"].number_format == "h:mm"
