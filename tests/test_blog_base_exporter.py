import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook
from openpyxl.utils.datetime import from_excel

from fet_to_xlsx.exporter.blog_base_exporter import BLOG_COLUMNS, BLOG_SHEET_NAME, generar_base_horarios_blog


def _time_value(value):
    if hasattr(value, "hour"):
        return value
    return from_excel(value).time()


def test_generar_base_horarios_blog_reads_actividades_and_preserves_other_sheets():
    workbook = Workbook()
    activities = workbook.active
    activities.title = "Actividades"
    activities.append([
        "ID", "Docente", "Materia", "Grupo", "Duración", "Cantidad de estudiantes",
        "Actividad activa", "Día", "Hora", "Salón", "Comentarios",
    ])
    activities.append([
        "1", "Ana", "Competencia", "2925097 - MANTENIMIENTO ELECTROMECANICO INDUSTRIAL . - MIXTA",
        4, 30, "Sí", "Lunes", "18:00", "Aula 1", "",
    ])
    activities.append(["2", "SIN DOCENTE", "Sin ubicar", "123 - Prog - DIURNA", 1, None, "Sí", "", "", "", ""])
    workbook.create_sheet("Horarios")["A1"] = "no tocar"
    workbook.create_sheet(BLOG_SHEET_NAME)["A1"] = "viejo"

    worksheet = generar_base_horarios_blog(workbook)

    assert workbook["Horarios"]["A1"].value == "no tocar"
    assert worksheet.title == BLOG_SHEET_NAME
    assert [cell.value for cell in worksheet[1]] == BLOG_COLUMNS
    assert worksheet.max_row == 2
    assert worksheet["A2"].value == "2925097"
    assert worksheet["B2"].value == "MANTENIMIENTO ELECTROMECANICO INDUSTRIAL"
    assert worksheet["C2"].value == "MIXTA"
    assert worksheet["D2"].value == "Lunes"
    assert _time_value(worksheet["E2"].value).hour == 18
    assert _time_value(worksheet["F2"].value).hour == 22
    assert worksheet["G2"].value == "Competencia"
    assert worksheet["H2"].value == "Ana"
    assert worksheet["I2"].value == "Aula 1"
    assert worksheet["J2"].value is None
    assert worksheet["K2"].value is None


def test_generar_base_horarios_blog_blanks_sin_docente():
    workbook = Workbook()
    activities = workbook.active
    activities.title = "Actividades"
    activities.append(["ID", "Docente", "Materia", "Grupo", "Duración", "Cantidad de estudiantes", "Actividad activa", "Día", "Hora", "Salón", "Comentarios"])
    activities.append(["1", "SIN DOCENTE", "Competencia", "Tco Mecánico 1 - Programa - TARDE", 1, None, "Sí", "Martes", "6:00", "", ""])

    worksheet = generar_base_horarios_blog(workbook)

    assert worksheet["A2"].value == "Tco Mecánico 1"
    assert worksheet["B2"].value == "Programa - TARDE"
    assert worksheet["C2"].value is None
    assert worksheet["H2"].value is None
